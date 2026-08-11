"""Eksport zakładek Raporty do jednego pliku Excel (openpyxl)."""
from accounts.models import Department
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .reports_views import _build_all_tabs, _report_filters

CLARET = "661C31"
CORAL = "FF4143"
GOLD_TINT = "FBF1DC"
STEEL_TINT = "E4ECF1"
CORAL_TINT = "FDEAEA"
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
HEADER_FILL = PatternFill("solid", fgColor=CLARET)
TITLE_FONT = Font(color=CLARET, bold=True, size=13, name="Calibri")
SUBTITLE_FONT = Font(color=CLARET, bold=True, size=11, name="Calibri")
BOLD = Font(bold=True, name="Calibri")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SCORE_FILLS = {
    "score-good": PatternFill("solid", fgColor=STEEL_TINT),
    "score-mid": PatternFill("solid", fgColor=GOLD_TINT),
    "score-bad": PatternFill("solid", fgColor=CORAL_TINT),
}


def _header_row(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, min_width=10, max_width=60):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = max(min_width, min(length + 2, max_width))


def _write_score_cell(ws, row, col, cell_dict):
    value = cell_dict.get("value") if isinstance(cell_dict, dict) else cell_dict
    css = cell_dict.get("css") if isinstance(cell_dict, dict) else ""
    c = ws.cell(row=row, column=col, value=f"{value}%" if value is not None else "—")
    c.border = BORDER
    c.alignment = Alignment(horizontal="center")
    if css in SCORE_FILLS:
        c.fill = SCORE_FILLS[css]
    return c


def _build_overview_sheet(wb, data, year, department_label):
    ws = wb.active
    ws.title = "Przeglad"
    ws["A1"] = f"Przeglad niezgodnosci {year}" + (f" - {department_label}" if department_label else "")
    ws["A1"].font = TITLE_FONT

    ws["A3"] = "Niezgodnosci w okresie"
    ws["B3"] = data["total_ncs"]
    ws["A4"] = "Przeterminowane"
    ws["B4"] = data["overdue_count"]
    ws["A5"] = "Sr. dni do wdrozenia dzialan"
    ws["B5"] = data["avg_days_to_close"] if data["avg_days_to_close"] is not None else "-"
    for r in (3, 4, 5):
        ws.cell(row=r, column=1).font = BOLD

    row = 7
    ws.cell(row=row, column=1, value="Status").font = SUBTITLE_FONT
    row += 1
    _header_row(ws, row, ["Status", "Liczba"])
    for entry in data["by_status"]:
        row += 1
        ws.cell(row=row, column=1, value=entry["status"]).border = BORDER
        ws.cell(row=row, column=2, value=entry["total"]).border = BORDER

    row += 2
    ws.cell(row=row, column=1, value="Top 10 lokalizacji").font = SUBTITLE_FONT
    row += 1
    _header_row(ws, row, ["Lokalizacja", "Liczba niezgodnosci"])
    for entry in data["top_locations"]:
        row += 1
        ws.cell(row=row, column=1, value=entry["location_detail"]).border = BORDER
        ws.cell(row=row, column=2, value=entry["total"]).border = BORDER

    _autosize(ws)


def _build_results_sheet(wb, data, year):
    ws = wb.create_sheet("Wyniki inspekcji")
    ws["A1"] = f"Wyniki inspekcji GMP/GHP {year}"
    ws["A1"].font = TITLE_FONT
    row = 3
    month_labels = data["month_labels"]

    for dept in data["results_by_department"]:
        ws.cell(row=row, column=1, value=dept["department_label"]).font = SUBTITLE_FONT
        row += 1
        if dept["has_score_data"]:
            _header_row(ws, row, ["Zmiana"] + month_labels + [f"Sr. {year}"])
            row += 1
            for shift_row in dept["shift_rows"]:
                ws.cell(row=row, column=1, value=shift_row["shift"]).font = BOLD
                ws.cell(row=row, column=1).border = BORDER
                for i, cell in enumerate(shift_row["months"]):
                    _write_score_cell(ws, row, 2 + i, cell)
                _write_score_cell(ws, row, 2 + len(month_labels), shift_row["year_avg"])
                row += 1
            ws.cell(row=row, column=1, value=f"Srednia dla dzialu ({year})").font = BOLD
            _write_score_cell(ws, row, 2 + len(month_labels), dept["department_avg"])
            row += 1
        else:
            ws.cell(row=row, column=1, value=f"Brak wynikow inspekcji w {year} r.")
            row += 1

        if dept["has_quarter_data"]:
            row += 1
            _header_row(ws, row, ["Rok", "Q1", "Q2", "Q3", "Q4", "Suma"])
            row += 1
            for qrow in dept["quarter_rows"]:
                ws.cell(row=row, column=1, value=qrow["year"]).font = BOLD
                ws.cell(row=row, column=1).border = BORDER
                for i, v in enumerate(qrow["q"]):
                    ws.cell(row=row, column=2 + i, value=v).border = BORDER
                ws.cell(row=row, column=6, value=qrow["total"]).font = BOLD
                row += 1
        row += 2

    _autosize(ws)


def _build_groups_sheet(wb, data, year):
    ws = wb.create_sheet("NC wg grup")
    ws["A1"] = f"Niezgodnosci wg grup GMP/GHP {year}"
    ws["A1"].font = TITLE_FONT
    row = 3

    for dept in data["groups_by_department"]:
        ws.cell(row=row, column=1, value=dept["department_label"]).font = SUBTITLE_FONT
        row += 1
        if dept["rows"]:
            _header_row(ws, row, ["Grupa niezgodnosci", "Q1", "Q2", "Q3", "Q4", "Suma", "Udzial %"])
            row += 1
            for r in dept["rows"]:
                ws.cell(row=row, column=1, value=r["category"]).border = BORDER
                for i, v in enumerate(r["q"]):
                    ws.cell(row=row, column=2 + i, value=v).border = BORDER
                ws.cell(row=row, column=6, value=r["total"]).font = BOLD
                ws.cell(row=row, column=7, value=f"{r['pct']}%").border = BORDER
                row += 1
            tr = dept["totals_row"]
            ws.cell(row=row, column=1, value="Suma").font = BOLD
            for i, v in enumerate(tr["q"]):
                ws.cell(row=row, column=2 + i, value=v).font = BOLD
            ws.cell(row=row, column=6, value=tr["total"]).font = BOLD
            row += 1
        else:
            ws.cell(row=row, column=1, value=f"Brak niezgodnosci w {year} r.")
            row += 1
        row += 1

    _autosize(ws)


def _build_points_sheet(wb, data, year):
    ws = wb.create_sheet("NC wg punktow")
    ws["A1"] = f"Niezgodnosci wg punktow checklisty {year}"
    ws["A1"].font = TITLE_FONT
    row = 3

    for dept in data["points_by_department"]:
        ws.cell(row=row, column=1, value=dept["department_label"]).font = SUBTITLE_FONT
        row += 1
        if dept["categories"]:
            for cat in dept["categories"]:
                ws.cell(row=row, column=1, value=f"{cat['category']} ({cat['category_total']})").font = BOLD
                row += 1
                _header_row(ws, row, ["Punkt", "Opis", "Q1", "Q2", "Q3", "Q4", "Suma"])
                row += 1
                for p in cat["points"]:
                    ws.cell(row=row, column=1, value=p["label"]).border = BORDER
                    ws.cell(row=row, column=2, value=p["description"]).border = BORDER
                    for i, v in enumerate(p["q"]):
                        ws.cell(row=row, column=3 + i, value=v).border = BORDER
                    ws.cell(row=row, column=7, value=p["total"]).font = BOLD
                    row += 1
                row += 1
        else:
            ws.cell(row=row, column=1, value=f"Brak niezgodnosci w {year} r.")
            row += 1
        row += 1

    _autosize(ws)


@login_required
def reports_export(request):
    if not request.user.can_view_reports:
        messages.error(request, "Raporty są dostępne tylko dla QualityAdmin i Helpdesku.")
        return redirect("dashboard")

    f = _report_filters(request)
    tabs = _build_all_tabs(f)
    department_label = dict(Department.choices).get(f["department"], "") if f["department"] else ""

    wb = Workbook()
    _build_overview_sheet(wb, tabs["przeglad"], f["year"], department_label)
    _build_results_sheet(wb, tabs["wyniki"], f["year"])
    _build_groups_sheet(wb, tabs["grupy"], f["year"])
    _build_points_sheet(wb, tabs["punkty"], f["year"])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    suffix = f"_{f['department']}" if f["department"] else ""
    response["Content-Disposition"] = f'attachment; filename="Raport_GMP_GHP_{f["year"]}{suffix}.xlsx"'
    wb.save(response)
    return response
