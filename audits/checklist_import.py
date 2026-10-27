"""
Import treści checklisty GMP/GHP z pliku Excel (wzór formularza CD-xxxx) do bazy:
ChecklistTemplate -> ChecklistSection -> ChecklistSubsection -> ChecklistItem.

Używany przez komendę `import_checklists_from_xlsx` oraz przez ekran importu
w aplikacji (tylko administrator) - dzięki temu treść formularza można wgrać
na produkcji bez wrzucania plików z realnymi danymi firmy do repozytorium.
"""
import re

from django.db import transaction

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from .models import ChecklistItem, ChecklistSection, ChecklistSubsection, ChecklistTemplate

ROMAN_RE = re.compile(r"^[IVXLC]+$")
# Nagłówek sekcji: cała treść w kolumnie A, np. "I. HIGIENA - PORZĄDEK" (kolumna B puste)
SECTION_MARKER_RE = re.compile(r"^([IVXLC]+)\.\s*(.+)$", re.DOTALL)
# Punkt kontrolny: czysty numer w A (np. "1.1", "1.1."), opis w kolumnie B
ITEM_RE = re.compile(r"^(\d+\.\d+)\.?$")
# Nagłówek podsekcji: numer + treść razem w kolumnie A, np. "1. Drobny sprzęt..." (kolumna B puste)
SUBSECTION_MARKER_RE = re.compile(r"^(\d+)\.\s*(.+)$", re.DOTALL)


class ChecklistImportError(Exception):
    """Plik nie wygląda jak wzór checklisty GMP/GHP - komunikat nadaje się do pokazania użytkownikowi."""


def guess_sheet_name(workbook):
    for ws in workbook.worksheets:
        if not ws.title.lower().startswith("arkusz"):
            return ws.title
    raise ChecklistImportError("Plik nie zawiera arkusza z checklistą.")


@transaction.atomic
def import_template(source, area_code, sheet_name=None, version_note="", existing="deactivate"):
    """Tworzy szablon checklisty z pliku Excel (ścieżka lub obiekt plikowy).

    existing: co zrobić z dotychczasowymi szablonami tego obszaru -
      "deactivate" (domyślnie, bezpieczne: stare zostają dla historii inspekcji) albo
      "delete" (kasuje je; nie powiedzie się, jeśli mają zapisane inspekcje).
    Przy błędzie nic nie zostaje zapisane (transakcja). Zwraca (szablon, liczba punktów)."""
    if openpyxl is None:
        raise ChecklistImportError("Pakiet openpyxl nie jest zainstalowany.")
    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception as exc:  # noqa: BLE001 - dowolny błąd odczytu pliku to dla użytkownika "zły plik"
        raise ChecklistImportError(f"Nie udało się odczytać pliku Excel: {exc}") from exc

    sheet_name = sheet_name or guess_sheet_name(wb)
    if sheet_name not in wb.sheetnames:
        raise ChecklistImportError(f"Brak arkusza „{sheet_name}” w pliku.")
    ws = wb[sheet_name]

    title = ws.cell(row=1, column=1).value or f"Inspekcja GMP/GHP - {area_code}"
    doc_code = ws.cell(row=2, column=1).value or ""
    max_row = ws.max_row

    # --- 1) tabela podsumowania: sekcje + max punktów -------------------------
    summary_header_row = None
    for r in range(1, max_row + 1):
        row_values = [ws.cell(row=r, column=c).value for c in range(1, 7)]
        if any(v and "zakres wymag" in str(v).lower() for v in row_values) and \
           any(v and "max ilo" in str(v).lower() for v in row_values):
            summary_header_row = r
            break
    if summary_header_row is None:
        raise ChecklistImportError("Nie znaleziono tabeli podsumowania (zakres wymagań / max ilość punktów).")

    previous = ChecklistTemplate.objects.filter(area_code=area_code)
    if existing == "delete":
        previous.delete()
    else:
        previous.update(is_active=False)

    template = ChecklistTemplate.objects.create(
        area_code=area_code,
        name=str(title).strip(),
        document_code=str(doc_code).strip(),
        version_note=version_note,
        is_active=True,
    )

    section_defs = []
    r = summary_header_row + 1
    order = 0
    while r <= max_row:
        label = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        max_points = ws.cell(row=r, column=4).value
        if label and ROMAN_RE.match(str(label).strip()):
            order += 1
            section_defs.append(ChecklistSection.objects.create(
                template=template, order=order, label=str(label).strip(),
                name=str(name).strip() if name else "",
                max_points=int(max_points) if max_points else 0,
            ))
            r += 1
        else:
            break
    if not section_defs:
        raise ChecklistImportError("Tabela podsumowania nie zawiera żadnych zakresów wymagań.")

    # --- 2) szczegółowa checklista: sekcje / podsekcje / punkty ---------------
    detail_header_row = None
    for r in range(summary_header_row, max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        if a and str(a).strip().lower() == "l.p." and b and "zakres wymag" in str(b).lower():
            detail_header_row = r
            break
    if detail_header_row is None:
        raise ChecklistImportError("Nie znaleziono nagłówka checklisty szczegółowej (l.p. / zakres wymagań).")

    current_section_idx = -1
    current_section = None
    current_subsection = None
    sub_order = 0
    item_order = 0
    items_count = 0

    for r in range(detail_header_row + 1, max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        a_str = str(a).strip() if a is not None else ""
        b_str = str(b).strip() if b is not None else ""

        if not a_str and not b_str:
            continue

        section_marker = SECTION_MARKER_RE.match(a_str) if a_str else None
        if section_marker and not b_str:
            current_section_idx += 1
            if current_section_idx < len(section_defs):
                current_section = section_defs[current_section_idx]
            current_subsection = None
            sub_order = 0
            item_order = 0
            continue

        if ITEM_RE.match(a_str) and current_subsection is not None:
            item_order += 1
            items_count += 1
            ChecklistItem.objects.create(
                subsection=current_subsection, order=item_order,
                number=a_str.rstrip("."), description=b_str,
            )
            continue

        sub_marker = SUBSECTION_MARKER_RE.match(a_str) if a_str else None
        if sub_marker and not b_str and current_section is not None:
            sub_order += 1
            item_order = 0
            current_subsection = ChecklistSubsection.objects.create(
                section=current_section, order=sub_order,
                number=sub_marker.group(1), name=sub_marker.group(2).strip(),
            )
            continue

    if not items_count:
        raise ChecklistImportError("Nie znaleziono żadnych punktów kontrolnych (numery typu 1.1, 2.3).")
    return template, items_count
