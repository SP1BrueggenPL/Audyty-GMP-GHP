"""
Jednorazowy import realnej treści checklist GMP/GHP z plików Excel
obowiązujących w zakładzie (folder Wytyczne/Templatki) do bazy danych.

Użycie:
    python manage.py import_checklists_from_xlsx

Skrypt czyta trzy "czyste" (niewypełnione) wzory formularzy i tworzy z nich
strukturę ChecklistTemplate -> ChecklistSection -> ChecklistSubsection ->
ChecklistItem. Jeśli szablon dla danego area_code już istnieje, jest
usuwany i tworzony od nowa (żeby komenda była bezpieczna do ponownego
uruchomienia po aktualizacji wzoru).
"""
import re
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from audits.models import AreaCode, ChecklistItem, ChecklistSection, ChecklistSubsection, ChecklistTemplate

ROMAN_RE = re.compile(r"^[IVXLC]+$")
# Nagłówek sekcji: cała treść w kolumnie A, np. "I. HIGIENA - PORZĄDEK" (kolumna B puste)
SECTION_MARKER_RE = re.compile(r"^([IVXLC]+)\.\s*(.+)$", re.DOTALL)
# Punkt kontrolny: czysty numer w A (np. "1.1", "1.1."), opis w kolumnie B
ITEM_RE = re.compile(r"^(\d+\.\d+)\.?$")
# Nagłówek podsekcji: numer + treść razem w kolumnie A, np. "1. Drobny sprzęt..." (kolumna B puste)
SUBSECTION_MARKER_RE = re.compile(r"^(\d+)\.\s*(.+)$", re.DOTALL)

SOURCE_DIR = settings.BASE_DIR.parent / "Wytyczne" / "Templatki"

TEMPLATES = [
    {
        "area_code": AreaCode.WCE_WPP,
        "file": "GMP_GHP_WPP_WCE_FORM_12.01.2026.xlsx",
        "sheet": "GMP GHP_WCE_WPP",
    },
    {
        "area_code": AreaCode.WLS,
        "file": "GMP_GHP_WLS_12.01.2026.xlsx",
        "sheet": "GMP GHP_WLS",
    },
    {
        "area_code": AreaCode.WED,
        "file": "GMP_WED_16.07.2026.xlsx",
        "sheet": "WED",
    },
]


class Command(BaseCommand):
    help = "Importuje strukturę checklist GMP/GHP z oryginalnych plików Excel do bazy danych."

    def add_arguments(self, parser):
        parser.add_argument(
            "--source-dir", default=None,
            help="Ścieżka do folderu z plikami xlsx (domyślnie ../Wytyczne/Templatki względem projektu).",
        )

    def handle(self, *args, **options):
        if openpyxl is None:
            raise CommandError("Pakiet openpyxl nie jest zainstalowany (pip install openpyxl).")

        source_dir = Path(options["source_dir"]) if options["source_dir"] else SOURCE_DIR
        if not source_dir.exists():
            raise CommandError(f"Nie znaleziono folderu ze wzorami: {source_dir}")

        for cfg in TEMPLATES:
            path = source_dir / cfg["file"]
            if not path.exists():
                self.stderr.write(self.style.WARNING(f"Pominięto {cfg['area_code']}: brak pliku {path}"))
                continue
            self.import_one(cfg, path)

    def import_one(self, cfg, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[cfg["sheet"]]

        title = ws.cell(row=1, column=1).value or f"Inspekcja GMP/GHP - {cfg['area_code']}"
        doc_code = ws.cell(row=2, column=1).value or ""

        ChecklistTemplate.objects.filter(area_code=cfg["area_code"]).delete()
        template = ChecklistTemplate.objects.create(
            area_code=cfg["area_code"],
            name=str(title).strip(),
            document_code=str(doc_code).strip(),
            version_note=path.stem,
            is_active=True,
        )

        max_row = ws.max_row

        # --- 1) tabela podsumowania: sekcje + max punktów -------------------
        summary_header_row = None
        for r in range(1, max_row + 1):
            row_values = [ws.cell(row=r, column=c).value for c in range(1, 7)]
            if any(v and "zakres wymag" in str(v).lower() for v in row_values) and \
               any(v and "max ilo" in str(v).lower() for v in row_values):
                summary_header_row = r
                break
        if summary_header_row is None:
            raise CommandError(f"Nie znaleziono tabeli podsumowania w {path.name}")

        section_defs = []  # (label, name, max_points) w kolejności
        r = summary_header_row + 1
        order = 0
        while r <= max_row:
            label = ws.cell(row=r, column=1).value
            name = ws.cell(row=r, column=2).value
            max_points = ws.cell(row=r, column=4).value
            if label and ROMAN_RE.match(str(label).strip()):
                order += 1
                section = ChecklistSection.objects.create(
                    template=template,
                    order=order,
                    label=str(label).strip(),
                    name=str(name).strip() if name else "",
                    max_points=int(max_points) if max_points else 0,
                )
                section_defs.append(section)
                r += 1
            else:
                break

        # --- 2) szczegółowa checklista: sekcje / podsekcje / punkty ---------
        detail_header_row = None
        for r in range(summary_header_row, max_row + 1):
            a = ws.cell(row=r, column=1).value
            b = ws.cell(row=r, column=2).value
            if a and str(a).strip().lower() == "l.p." and b and "zakres wymag" in str(b).lower():
                detail_header_row = r
                break
        if detail_header_row is None:
            raise CommandError(f"Nie znaleziono nagłówka checklisty szczegółowej w {path.name}")

        current_section_idx = -1
        current_section = None
        current_subsection = None
        sub_order = 0
        item_order = 0

        for r in range(detail_header_row + 1, max_row + 1):
            a = ws.cell(row=r, column=1).value
            b = ws.cell(row=r, column=2).value
            a_str = str(a).strip() if a is not None else ""
            b_str = str(b).strip() if b is not None else ""

            if not a_str and not b_str:
                continue

            section_marker = SECTION_MARKER_RE.match(a_str) if a_str else None
            if section_marker and not b_str:
                current_section_idx += 1
                if current_section_idx < len(section_defs):
                    current_section = section_defs[current_section_idx]
                current_subsection = None
                sub_order = 0
                item_order = 0
                continue

            if ITEM_RE.match(a_str) and current_subsection is not None:
                item_order += 1
                ChecklistItem.objects.create(
                    subsection=current_subsection,
                    order=item_order,
                    number=a_str.rstrip("."),
                    description=b_str,
                )
                continue

            sub_marker = SUBSECTION_MARKER_RE.match(a_str) if a_str else None
            if sub_marker and not b_str and current_section is not None:
                sub_order += 1
                item_order = 0
                current_subsection = ChecklistSubsection.objects.create(
                    section=current_section,
                    order=sub_order,
                    number=sub_marker.group(1),
                    name=sub_marker.group(2).strip(),
                )
                continue

        self.stdout.write(self.style.SUCCESS(
            f"Zaimportowano {cfg['area_code']}: {len(section_defs)} sekcji z {path.name}"
        ))
